from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.parts.image import ImagePart
nsmap['asvg'] = 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'
from PIL import Image, ImageDraw, ImageFont
import os
import subprocess
import tempfile
from io import BytesIO
try:
    import cairosvg
except Exception:
    cairosvg = None
try:
    from svglib.svglib import svg2rlg
    from reportlab.graphics import renderPDF
except Exception:
    svg2rlg = None
    renderPDF = None

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(os.path.dirname(SCRIPT_DIR))
OUT = os.path.join(PROJECT_ROOT, "02_reference", "proposals", "NEDO_企画書_衛星データ活用3案.docx")
FONT = "Hiragino Sans"

NAVY = "0B2545"
BLUE = "1F4D78"
MID = "5B6B7A"
LIGHT = "EAF0F6"
PALE = "F5F7FA"
GOLD = "8A6500"
RED = "9B1C1C"
WHITE = "FFFFFF"

ASSET_DIR = os.path.join(SCRIPT_DIR, "assets", "proposal_assets")
ICON_DIR = os.path.join(SCRIPT_DIR, "assets", "external_icons", "tabler")

def make_diagram(kind, path):
    """Create a compact, vector-like visual explainer for the one-pager."""
    W, H = 1500, 170
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    font_path = "/System/Library/Fonts/ヒラギノ角ゴシック W4.ttc"
    font_bold_path = "/System/Library/Fonts/ヒラギノ角ゴシック W8.ttc"
    try:
        f_label = ImageFont.truetype(font_path, 29)
        f_small = ImageFont.truetype(font_path, 22)
        f_bold = ImageFont.truetype(font_bold_path, 30)
    except Exception:
        f_label = f_small = f_bold = ImageFont.load_default()
    navy = "#0B2545"; blue = "#1F4D78"; pale = "#EAF3F8"; line = "#9AB2C8"; gold = "#C38B18"
    boxes = [(22, 25, 330, 145), (410, 25, 718, 145), (798, 25, 1106, 145), (1186, 25, 1494, 145)]
    labels = {
        1: ["衛星観測", "砂州・樹木", "測量候補", "掘削・伐採"],
        2: ["衛星観測", "港湾・道路", "残存機能", "代替輸送"],
        3: ["衛星観測", "被害・通行", "使える拠点", "支援ルート"],
    }[kind]
    icon_names = {
        1: ["satellite", "trees", "map", "road"],
        2: ["satellite", "ship", "map", "route"],
        3: ["satellite", "road", "map", "truck"],
    }[kind]
    descs = {
        1: ["河道を反復観測", "流下を妨げる変化", "現地確認を優先", "予算・班を配分"],
        2: ["浸水・堆積を把握", "岸壁＋背後地", "港湾の残存能力", "港・陸路を切替"],
        3: ["広域の変化を把握", "閉塞・孤立を確認", "信頼度付き更新", "物資・燃料を投入"],
    }[kind]
    for i, (x1, y1, x2, y2) in enumerate(boxes):
        fill = pale if i in (0, 3) else "#F7F9FB"
        d.rounded_rectangle((x1, y1, x2, y2), radius=18, fill=fill, outline=line, width=3)
        # simple icons, deliberately pictographic rather than decorative
        cx = (x1 + x2) // 2; cy = 62
        if i == 0:  # satellite + target-specific observation cue
            d.rectangle((cx-25, cy-12, cx+25, cy+12), fill=blue, outline=navy, width=3)
            d.line((cx-25, cy, cx-62, cy-25), fill=navy, width=5)
            d.line((cx+25, cy, cx+62, cy-25), fill=navy, width=5)
            d.rectangle((cx-82, cy-42, cx-60, cy-20), fill=gold, outline=navy, width=2)
            d.rectangle((cx+60, cy-42, cx+82, cy-20), fill=gold, outline=navy, width=2)
            d.line((cx-80, cy+47, cx+80, cy+47), fill=gold, width=3)
        elif i == 1 and kind == 1:  # river corridor, sandbar and trees
            d.polygon([(cx-78, cy-35), (cx-20, cy-12), (cx+20, cy-28), (cx+78, cy-8), (cx+78, cy+22), (cx+20, cy+6), (cx-20, cy+24), (cx-78, cy+5)], fill="#A7D2E8", outline=blue)
            d.polygon([(cx-42, cy-2), (cx-5, cy-10), (cx+20, cy+1), (cx-10, cy+13)], fill="#D0A15D", outline=gold)
            for tx in (38, 60):
                d.line((cx+tx, cy+8, cx+tx, cy-22), fill="#6A7F52", width=4)
                d.ellipse((cx+tx-13, cy-37, cx+tx+13, cy-12), fill="#6A9A59", outline=navy, width=2)
        elif i == 1 and kind == 2:  # port, ship, crane and access road
            d.line((cx-78, cy+30, cx+78, cy+30), fill=blue, width=6)
            d.rectangle((cx-62, cy-15, cx-8, cy+18), fill="#C8D5DF", outline=navy, width=2)
            d.line((cx-42, cy-15, cx-42, cy-47), fill=navy, width=5); d.line((cx-42, cy-47, cx-5, cy-15), fill=navy, width=4)
            d.polygon([(cx+15, cy+10), (cx+62, cy+10), (cx+48, cy+25), (cx+25, cy+25)], fill=gold, outline=navy)
            d.line((cx+78, cy+30, cx+78, cy+52), fill=navy, width=4); d.line((cx+78, cy+52, cx+25, cy+52), fill=navy, width=4)
        elif i == 1 and kind == 3:  # damaged road and isolated hub
            d.line((cx-78, cy+28, cx-30, cy+3), fill="#C94C4C", width=8)
            d.line((cx-8, cy-7, cx+30, cy-30), fill="#C94C4C", width=8)
            d.line((cx+30, cy-30, cx+78, cy-4), fill="#6BA56B", width=8)
            d.ellipse((cx-43, cy-10, cx-17, cy+16), fill="#D85B5B", outline=navy, width=2)
            d.ellipse((cx+62, cy-18, cx+88, cy+8), fill="#69A76C", outline=navy, width=2)
            d.rectangle((cx+7, cy+5, cx+35, cy+32), fill=gold, outline=navy, width=2)
        elif i == 2 and kind == 1:  # survey pin and measurement strip
            d.polygon([(cx, cy-44), (cx-19, cy-15), (cx, cy+11), (cx+19, cy-15)], fill=gold, outline=navy)
            d.ellipse((cx-7, cy-23, cx+7, cy-9), fill="#FFFFFF", outline=navy)
            d.line((cx-74, cy+35, cx+75, cy+35), fill=blue, width=5)
            for px in (-45, -15, 15, 45): d.line((cx+px, cy+25, cx+px, cy+45), fill=blue, width=3)
        elif i == 2 and kind == 2:  # status score for port assets
            d.rectangle((cx-65, cy-35, cx-30, cy+35), fill="#72A9C9", outline=navy, width=2)
            d.rectangle((cx-12, cy-5, cx+23, cy+35), fill=gold, outline=navy, width=2)
            d.rectangle((cx+41, cy-50, cx+76, cy+35), fill="#D98383", outline=navy, width=2)
            d.line((cx-80, cy+45, cx+82, cy+45), fill=navy, width=3)
        elif i == 2 and kind == 3:  # green usable hubs and confidence ring
            d.ellipse((cx-40, cy-40, cx+40, cy+40), outline=blue, width=6)
            d.ellipse((cx-20, cy-20, cx+20, cy+20), fill="#69A76C", outline=navy, width=2)
            d.line((cx-85, cy+35, cx-42, cy+7), fill="#69A76C", width=6)
            d.line((cx+42, cy+7, cx+85, cy-30), fill="#69A76C", width=6)
        elif i == 2:  # map / confidence
            d.polygon([(cx-68, cy-28), (cx-25, cy-43), (cx+20, cy-25), (cx+67, cy-40), (cx+67, cy+35), (cx+20, cy+20), (cx-25, cy+38), (cx-68, cy+22)], fill="#FFFFFF", outline=blue)
            d.line((cx-25, cy-43, cx-25, cy+38), fill=line, width=3)
            d.line((cx+20, cy-25, cx+20, cy+20), fill=line, width=3)
            d.ellipse((cx-10, cy-4, cx+10, cy+16), fill=gold, outline=navy, width=2)
        elif i == 3 and kind == 1:  # maintenance actions
            d.rectangle((cx-60, cy-15, cx+25, cy+12), fill="#C8D5DF", outline=navy, width=2)
            d.rectangle((cx+25, cy-32, cx+56, cy+12), fill=gold, outline=navy, width=2)
            d.ellipse((cx-45, cy+10, cx-15, cy+40), fill=navy)
            d.ellipse((cx+33, cy+10, cx+63, cy+40), fill=navy)
            d.line((cx+72, cy-45, cx+72, cy+15), fill=navy, width=4)
            d.line((cx+45, cy-12, cx+72, cy-45), fill=navy, width=4)
        elif i == 3 and kind == 2:  # alternative route network
            pts = [(cx-76, cy+25), (cx-25, cy-30), (cx+20, cy+22), (cx+75, cy-28)]
            for a, b in zip(pts, pts[1:]): d.line((a, b), fill="#C94C4C", width=5)
            d.line((cx-76, cy+25, cx+20, cy+22), fill="#69A76C", width=7)
            for px, py in pts: d.ellipse((px-11, py-11, px+11, py+11), fill=gold, outline=navy, width=2)
        elif i == 3 and kind == 3:  # relief truck and route
            d.line((cx-75, cy+30, cx+65, cy-30), fill="#69A76C", width=8)
            d.rectangle((cx-38, cy-15, cx+30, cy+18), fill=gold, outline=navy, width=2)
            d.rectangle((cx+30, cy-3, cx+58, cy+18), fill="#D9E3EB", outline=navy, width=2)
            d.ellipse((cx-23, cy+11, cx+3, cy+37), fill=navy)
            d.ellipse((cx+34, cy+11, cx+60, cy+37), fill=navy)
        else:  # generic network fallback
            pts = [(cx-58, cy+22), (cx-8, cy-30), (cx+45, cy+25), (cx+68, cy-35)]
            for a, b in zip(pts, pts[1:]): d.line((a, b), fill=blue, width=6)
            for px, py in pts: d.ellipse((px-11, py-11, px+11, py+11), fill=gold, outline=navy, width=2)
        # Use finished MIT-licensed Tabler SVGs when CairoSVG is available.
        # The hand-drawn fallback remains for environments without SVG support.
        if cairosvg is not None:
            icon_file = os.path.join(ICON_DIR, icon_names[i] + ".svg")
            if os.path.exists(icon_file):
                svg = open(icon_file, "r", encoding="utf-8").read().replace("currentColor", "#1F4D78")
                png = cairosvg.svg2png(bytestring=svg.encode("utf-8"), output_width=82, output_height=82)
                icon = Image.open(BytesIO(png)).convert("RGBA")
                img.paste(icon, (cx - 41, cy - 42), icon)
        bbox = d.textbbox((0, 0), labels[i], font=f_bold)
        d.text((cx - (bbox[2]-bbox[0])/2, 92), labels[i], font=f_bold, fill=navy)
        bbox2 = d.textbbox((0, 0), descs[i], font=f_small)
        d.text((cx - (bbox2[2]-bbox2[0])/2, 128), descs[i], font=f_small, fill=blue)
        if i < 3:
            d.line((x2+18, 85, x2+70, 85), fill=blue, width=5)
            d.polygon([(x2+70, 85), (x2+53, 75), (x2+53, 95)], fill=blue)
    img.save(path, "PNG", optimize=True)

def add_diagram(doc, kind):
    os.makedirs(ASSET_DIR, exist_ok=True)
    svg_path = os.path.join(ASSET_DIR, f"diagram_{kind}.svg")
    png_path = os.path.join(ASSET_DIR, f"diagram_{kind}.png")
    make_composite_svg(kind, svg_path)
    make_official_icon_diagram(kind, png_path)
    p = doc.add_paragraph()
    set_para(p, before=0, after=3, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    # Display a Word-compatible render of the official SVG assets.
    p.add_run().add_picture(png_path, width=Inches(6.45))

def make_official_icon_diagram(kind, path):
    """Compose the one-pager using the downloaded Tabler SVGs as the primary icons."""
    W, H = 1500, 300
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    font_path = "/System/Library/Fonts/ヒラギノ角ゴシック W4.ttc"
    font_bold_path = "/System/Library/Fonts/ヒラギノ角ゴシック W8.ttc"
    f_small = ImageFont.truetype(font_path, 25)
    f_bold = ImageFont.truetype(font_bold_path, 38)
    labels = {
        1: [("衛星観測", "河道を反復観測", "satellite"), ("砂州・樹木", "流下を妨げる変化", "trees"), ("測量候補", "現地確認を優先", "map"), ("掘削・伐採", "予算・班を配分", "road")],
        2: [("衛星観測", "浸水・堆積を把握", "satellite"), ("港湾・道路", "岸壁＋背後地", "ship"), ("残存機能", "港湾の残存能力", "map"), ("代替輸送", "港・陸路を切替", "route")],
        3: [("衛星観測", "広域の変化を把握", "satellite"), ("被害・通行", "閉塞・孤立を確認", "road"), ("使える拠点", "信頼度付き更新", "map") , ("支援ルート", "物資・燃料を投入", "truck")],
    }[kind]
    for i, (label, desc, icon_name) in enumerate(labels):
        x1 = 22 + i * 388; x2 = x1 + 308
        d.rounded_rectangle((x1, 18, x2, 282), radius=22, fill="#EAF3F8" if i in (0, 3) else "#F7F9FB", outline="#9AB2C8", width=4)
        icon_file = os.path.join(ICON_DIR, icon_name + ".svg")
        svg = open(icon_file, "r", encoding="utf-8").read().replace("currentColor", "#1F4D78")
        if cairosvg is not None:
            icon_png = cairosvg.svg2png(bytestring=svg.encode("utf-8"), output_width=116, output_height=116)
            icon = Image.open(BytesIO(icon_png)).convert("RGBA")
        else:
            with tempfile.TemporaryDirectory(prefix="tabler_render_") as td:
                tmp_svg = os.path.join(td, "icon.svg")
                tmp_pdf = os.path.join(td, "icon.pdf")
                tmp_png_base = os.path.join(td, "icon")
                with open(tmp_svg, "w", encoding="utf-8") as fh: fh.write(svg)
                drawing = svg2rlg(tmp_svg)
                renderPDF.drawToFile(drawing, tmp_pdf)
                subprocess.run(["pdftoppm", "-png", "-singlefile", "-r", "144", tmp_pdf, tmp_png_base], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
                icon = Image.open(tmp_png_base + ".png").convert("RGBA")
                icon.thumbnail((116, 116), Image.Resampling.LANCZOS)
        img.paste(icon, (x1 + 96, 35), icon)
        bbox = d.textbbox((0, 0), label, font=f_bold)
        d.text((x1 + 154 - (bbox[2]-bbox[0])/2, 165), label, font=f_bold, fill="#0B2545")
        bbox = d.textbbox((0, 0), desc, font=f_small)
        d.text((x1 + 154 - (bbox[2]-bbox[0])/2, 220), desc, font=f_small, fill="#1F4D78")
        if i < 3:
            ax = x2 + 18
            d.line((ax, 150, ax + 48, 150), fill="#1F4D78", width=7)
            d.polygon([(ax + 48, 150), (ax + 30, 138), (ax + 30, 162)], fill="#1F4D78")
    img.save(path, "PNG", optimize=True)

def make_composite_svg(kind, path):
    labels = {
        1: [("衛星観測", "河道を反復観測", "satellite"), ("砂州・樹木", "流下を妨げる変化", "trees"), ("測量候補", "現地確認を優先", "map"), ("掘削・伐採", "予算・班を配分", "road")],
        2: [("衛星観測", "浸水・堆積を把握", "satellite"), ("港湾・道路", "岸壁＋背後地", "ship"), ("残存機能", "港湾の残存能力", "map"), ("代替輸送", "港・陸路を切替", "route")],
        3: [("衛星観測", "広域の変化を把握", "satellite"), ("被害・通行", "閉塞・孤立を確認", "road"), ("使える拠点", "信頼度付き更新", "map"), ("支援ルート", "物資・燃料を投入", "truck")],
    }[kind]
    parts = []
    for _, _, icon in labels:
        raw = open(os.path.join(ICON_DIR, icon + ".svg"), encoding="utf-8").read()
        inner = raw[raw.find(">") + 1:raw.rfind("</svg>")]
        parts.append(inner.replace("currentColor", "#1F4D78"))
    svg = ['<svg xmlns="http://www.w3.org/2000/svg" width="1500" height="180" viewBox="0 0 1500 180">',
           '<style>text{font-family:"Hiragino Sans","Noto Sans CJK JP",sans-serif;fill:#0B2545} .small{font-size:22px;fill:#1F4D78} .label{font-size:30px;font-weight:700}</style>']
    for i, ((label, desc, _), inner) in enumerate(zip(labels, parts)):
        x = 22 + i * 388
        fill = '#EAF3F8' if i in (0, 3) else '#F7F9FB'
        svg.append(f'<rect x="{x}" y="18" width="308" height="142" rx="18" fill="{fill}" stroke="#9AB2C8" stroke-width="3"/>')
        svg.append(f'<g transform="translate({x+113} 30) scale(3.4)">{inner}</g>')
        svg.append(f'<text x="{x+154}" y="112" text-anchor="middle" class="label">{label}</text>')
        svg.append(f'<text x="{x+154}" y="140" text-anchor="middle" class="small">{desc}</text>')
        if i < 3:
            ax = x + 326
            svg.append(f'<path d="M{ax} 89h48l-14-9m14 9-14 9" fill="none" stroke="#1F4D78" stroke-width="5" stroke-linecap="round" stroke-linejoin="round"/>')
    svg.append('</svg>')
    with open(path, 'w', encoding='utf-8') as f:
        f.write(''.join(svg))

def attach_svg_asset(doc, svg_path):
    """Keep the original SVG as an embedded package asset for Office/editing workflows."""
    package = doc.part.package
    svg_partname = package.next_partname('/word/media/diagram%d.svg')
    svg_part = ImagePart(svg_partname, 'image/svg+xml', open(svg_path, 'rb').read())
    doc.part.relate_to(svg_part, RT.IMAGE)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=70, start=100, bottom=70, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_cell_border(cell, color="D4DCE5", sz="6"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), sz)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)

def set_run(run, size=8.7, color=NAVY, bold=False, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:ascii'), FONT)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), FONT)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic

def set_para(p, before=0, after=3, line=1.05, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align

def add_text(p, text, size=8.7, color=NAVY, bold=False, italic=False):
    r = p.add_run(text)
    set_run(r, size, color, bold, italic)
    return r

def add_label_para(doc, label, text, after=3):
    p = doc.add_paragraph()
    set_para(p, after=after)
    add_text(p, label + "  ", 8.7, BLUE, True)
    add_text(p, text, 8.7, NAVY)
    return p

def add_bullets(cell, items, size=8.1):
    for item in items:
        p = cell.add_paragraph(style='List Bullet')
        set_para(p, after=1, line=1.0)
        add_text(p, item, size, NAVY)

def add_box(doc, title, body, fill=PALE, title_color=BLUE, body_size=8.35):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell = t.cell(0, 0)
    cell.width = Inches(6.45)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, fill)
    set_cell_border(cell, "D4DCE5", "5")
    set_cell_margins(cell, 75, 130, 75, 130)
    p = cell.paragraphs[0]
    set_para(p, after=2, line=1.0)
    add_text(p, title, 8.5, title_color, True)
    p2 = cell.add_paragraph()
    set_para(p2, after=0, line=1.0)
    add_text(p2, body, body_size, NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def add_two_col(doc, left_title, left_items, right_title, right_items):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [Inches(3.17), Inches(3.17)]
    for i, cell in enumerate(t.rows[0].cells):
        cell.width = widths[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_shading(cell, WHITE)
        set_cell_border(cell, "D4DCE5", "5")
        set_cell_margins(cell, 65, 105, 55, 105)
    for cell, title, items in [(t.cell(0,0), left_title, left_items), (t.cell(0,1), right_title, right_items)]:
        p = cell.paragraphs[0]
        set_para(p, after=2, line=1.0)
        add_text(p, title, 8.3, BLUE, True)
        add_bullets(cell, items, 7.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def add_kpi_table(doc, rows):
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [Inches(2.1), Inches(2.1), Inches(2.1)]
    headers = ["評価軸", "プロトタイプで示すこと", "目安となる指標"]
    for i, cell in enumerate(t.rows[0].cells):
        cell.width = widths[i]
        set_cell_shading(cell, LIGHT)
        set_cell_border(cell, "C6D2DE", "5")
        set_cell_margins(cell, 55, 90, 55, 90)
        p = cell.paragraphs[0]
        set_para(p, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_text(p, headers[i], 7.7, BLUE, True)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = widths[i]
            set_cell_border(cells[i], "D4DCE5", "5")
            set_cell_margins(cells[i], 55, 90, 55, 90)
            p = cells[i].paragraphs[0]
            set_para(p, after=0, line=1.0)
            add_text(p, value, 7.45, NAVY)
    return t

def add_page(doc, proposal):
    if len(doc.paragraphs) > 1:
        doc.add_page_break()
    p = doc.add_paragraph()
    set_para(p, after=1, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, "NEDO Challenge, Satellite Data  |  1ページ企画書", 7.8, MID, True)
    p = doc.add_paragraph()
    set_para(p, after=1, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, proposal['title'], 17, NAVY, True)
    p = doc.add_paragraph()
    set_para(p, after=5, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, proposal['subtitle'], 9.2, MID)

    t = doc.add_table(rows=1, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    meta = [("テーマ", proposal['theme']), ("主利用者", proposal['user']), ("衛星の役割", proposal['sat_role']), ("GRIDの役割", proposal['grid_role'])]
    for i, (lab, val) in enumerate(meta):
        cell = t.cell(0, i)
        cell.width = Inches(1.58)
        set_cell_shading(cell, LIGHT)
        set_cell_border(cell, "C6D2DE", "5")
        set_cell_margins(cell, 45, 65, 45, 65)
        p = cell.paragraphs[0]
        set_para(p, after=1, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_text(p, lab, 7.1, BLUE, True)
        p2 = cell.add_paragraph()
        set_para(p2, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_text(p2, val, 7.0, NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

    add_diagram(doc, proposal['diagram'])

    add_box(doc, "提案の核心", proposal['core'], fill="EAF3F8", title_color=BLUE, body_size=8.7)
    add_two_col(doc, "社会課題", proposal['problem'], "衛星だからできること", proposal['satellite'])
    add_two_col(doc, "サービス像", proposal['service'], "GRIDを使う理由", proposal['grid'])
    add_kpi_table(doc, proposal['kpis'])
    p = doc.add_paragraph()
    set_para(p, before=4, after=2, line=1.0)
    add_text(p, "最大の反論と回答  ", 8.1, GOLD, True)
    add_text(p, proposal['risk'], 8.1, NAVY)
    p = doc.add_paragraph()
    set_para(p, before=2, after=0, line=1.0)
    add_text(p, "出典：", 6.6, MID, True)
    add_text(p, proposal['sources'] + '｜図版アイコン：Tabler Icons（MIT）', 6.6, MID)

def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.62)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.3)

    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn('w:ascii'), FONT)
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), FONT)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    normal.font.size = Pt(8.7)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.05
    for style_name in ['List Bullet']:
        s = doc.styles[style_name]
        s.font.name = FONT
        s._element.rPr.rFonts.set(qn('w:ascii'), FONT)
        s._element.rPr.rFonts.set(qn('w:hAnsi'), FONT)
        s._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        s.font.size = Pt(7.8)
        s.paragraph_format.left_indent = Inches(0.18)
        s.paragraph_format.first_line_indent = Inches(-0.12)

    footer = sec.footer.paragraphs[0]
    set_para(footer, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_text(footer, "GRID検討用ドラフト  |  2026年7月", 6.5, MID)

    proposals = [
        {
            'title': '河道の「流下能力変化」モニタリング',
            'diagram': 1,
            'subtitle': '河川の広域・反復観測から、測量・掘削・伐採が必要な区間を発見する',
            'theme': 'テーマ1', 'user': '河川管理者', 'sat_role': '河道・河岸の時系列変化', 'grid_role': '保全計画',
            'core': '衛星で河道内の土砂堆積・砂州・植生・河岸変化を追跡し、現地測量・河道掘削・樹木伐採の候補区間を抽出する。衛星は水深や安全性を断定せず、広域スクリーニングを担う。',
            'problem': ['河道延長が長く、現地測量を高頻度に実施できない', '土砂・樹木・河岸変化が流下能力や洪水リスクに影響する', '管理者が次に調査すべき区間を絞りにくい'],
            'satellite': ['SAR・光学・標高・懸濁物質データを時系列統合', '砂州・濁水・植生・河岸線の変化候補を抽出', '広域の変化を同じ基準で比較'],
            'service': ['河道区間ごとの変化・信頼度マップ', '測量・掘削・伐採候補のリスト', '現地確認後に状態を更新する管理画面'],
            'grid': ['区間別の保全優先順位', '予算・施工班・環境制約を含む複数年度計画', '流下能力低下リスクと維持費の比較'],
            'kpis': [('検出', '過去の測量・点検記録と照合', '変化候補の再現率'), ('業務', '候補区間を現地担当者が確認', '測量対象の絞り込み率'), ('価値', '衛星なしケースと比較', '調査時間・費用の削減')],
            'risk': '衛星から河床高や航路安全性を直接判定できない。対策として、出力を「測量・点検候補」に限定し、音響測深等の現地確認を前提にする。',
            'sources': 'NEDO公募・提供データ、国交省「河川維持管理基準」'
        },
        {
            'title': '臨海部の物流・産業機能低下モニタリング',
            'diagram': 2,
            'subtitle': '港湾単体ではなく、背後地・道路・物流拠点を含む「残存機能」を広域把握する',
            'theme': 'テーマ2', 'user': '港湾管理者・物流事業者', 'sat_role': '臨海部の面的な状態変化', 'grid_role': '代替シナリオ',
            'core': '港湾・ヤード・アクセス道路・工場・燃料基地などの被災・機能低下候補を衛星で把握し、港湾機能がどこまで残っているかを信頼度付きで提示する。安全性や荷役能力を衛星だけで断定しない。',
            'problem': ['港湾災害は、岸壁だけでなく背後地・道路・電力・物流の連鎖で影響が広がる', '複数組織の情報が分散し、臨海部全体の状況が見えにくい', '現地に入れない時間帯の初期把握が遅れる'],
            'satellite': ['SAR・光学で浸水・堆積・漂流物・アクセス障害候補を把握', '複数港・背後地を同じ基準で比較', '災害前後の面的変化を継続更新'],
            'service': ['臨海部の機能低下・残存機能マップ', '港湾・道路・物流拠点の信頼度付き状態', '追加確認が必要な地点と情報不足の明示'],
            'grid': ['港湾機能の部分低下をシナリオ化', '代替港・代替陸上輸送の比較', '貨物滞留・切替影響の推定'],
            'kpis': [('把握', '過去災害・荒天事例で再現', '初期状況把握時間'), ('衛星価値', '公開情報のみと比較', '追加検出された陸上障害'), ('実用', '担当者レビューで評価', 'シナリオの採用可能率')],
            'risk': 'サイバーポート等との重複、港湾安全性の断定、船社データ不足がリスク。主利用者を港湾管理者か船社のどちらかに絞り、衛星は陸上側の不完全情報に限定する。',
            'sources': 'NEDO公募・提供データ、国交省・JAXAの港湾衛星活用、GRID海運事例'
        },
        {
            'title': '災害後の「使えるインフラ」残存能力マップ',
            'diagram': 3,
            'subtitle': '被害箇所の列挙から、残っている社会機能の把握へ',
            'theme': 'テーマ2', 'user': '災害対策本部・応援調整部門', 'sat_role': '残存状態の広域推定', 'grid_role': '支援シナリオ',
            'core': '被災箇所を探すだけでなく、道路・港湾・物流拠点・避難拠点などの「利用可能性」を衛星・地上情報から信頼度付きで更新する。衛星は稼働を断定せず、支援判断の候補を作る。',
            'problem': ['災害初期は被害情報が断片的で、使える拠点や経路が分からない', '設備単位の被害情報を重ねても、社会機能の残存状況が見えにくい', '支援物資・燃料・要員の投入先を決める材料が不足する'],
            'satellite': ['光学・SAR・夜間光・土地利用変化を統合', '浸水・閉塞・周辺変化のない区域を広域抽出', '被害情報と残存候補を同じ地図上で更新'],
            'service': ['道路・港湾・拠点ごとの利用可能性と信頼度', '追加確認が必要な地点', '災害前後の残存機能の変化履歴'],
            'grid': ['残存する拠点を結ぶ代替ネットワーク分析', '支援物資・燃料・要員の輸送候補比較', '新情報による支援シナリオ更新'],
            'kpis': [('推定', '過去災害の時系列で検証', '残存候補の捕捉率'), ('判断', '被害マップのみと比較', '支援先候補の確定時間'), ('信頼度', '地上報告と照合', '誤判定・未確認率')],
            'risk': '物理的に残っていても稼働しているとは限らない。出力は「利用可能性」「追加確認要否」「信頼度」とし、通信・電力・人員などの地上情報と組み合わせる。',
            'sources': 'NEDO公募・提供データ、国交省DiMAPS・PLATEAU等、GRIDの物流・エネルギー領域'
        }
    ]
    for proposal in proposals:
        add_page(doc, proposal)
    doc.save(OUT)

if __name__ == '__main__':
    build()
